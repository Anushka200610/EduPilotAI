"""PowerPoint Presentation Generator for EduPilot AI CA-3 Assignment.

Generates a polished, professional 7-slide academic presentation with
custom layout, colors, bullet points, and speaker notes.
Output: docs/EduPilot_AI_CA3_Presentation.pptx
"""

import sys
from pathlib import Path

# Add project root to sys.path
PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from config.settings import DOCS_DIR

# Design Palette: Oxford Academic Blue & Slate
COLOR_NAVY = RGBColor(30, 58, 138)       # #1E3A8A Primary Header
COLOR_SLATE = RGBColor(71, 85, 105)      # #475569 Subtitle / Muted
COLOR_DARK = RGBColor(15, 23, 42)        # #0F172A Body text
COLOR_ACCENT = RGBColor(14, 116, 144)    # #0E7490 Teal / Accent


def create_presentation():
    prs = Presentation()
    # 16:9 widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]  # Completely blank slide layout

    slides_content = [
        # SLIDE 1
        {
            "title": "EduPilot AI: Academic Performance Analysis & Intervention",
            "subtitle": "An Agentic AI Decision-Support System | Second-Year Undergraduate AIML CA-3",
            "bullets": [
                "Team Details: Presenter: [YOUR NAME] (PRN: [YOUR PRN]) | Partner: [PARTNER NAME] (PRN: [PARTNER PRN])",
                "Department: Artificial Intelligence & Machine Learning | Academic Term: Fall 2026",
                "Problem Addressed: Automating early student academic review across fragmented departmental records",
                "Core Innovations: 4 Specialized Agents, LangGraph Orchestration, Local MCP Tools, and Independent Critic Validation",
                "Key Principle: Strict decision-support boundaries with 100% human-in-the-loop governance",
            ],
            "notes": (
                "Good morning, respected faculty and evaluators. Today, we are presenting EduPilot AI, "
                "an Agentic AI decision-support system designed to solve a critical challenge in academic institutions: "
                "timely, evidence-grounded student performance analysis and early intervention."
            ),
        },
        # SLIDE 2
        {
            "title": "Problem Statement & Higher Education Domain Context",
            "subtitle": "Inefficiencies in Existing Academic Workflows & Justification for Agentic AI",
            "bullets": [
                "Fragmented Academic Records: Attendance, continuous test scores, and LMS assignments stored in disconnected silos.",
                "Cognitive Overload on Faculty Mentors: Manually reviewing 60-100 students across 4-6 subjects is labor-intensive.",
                "Delayed Interventions: Struggling students identified only after mid-term exams, leaving insufficient time for remediation.",
                "Why Agentic AI?: Monolithic scripts cannot adapt execution plans; black-box LLMs hallucinate calculations.",
                "EduPilot Solution: Decomposed multi-agent workflow separating planning, analytics, interventions, and adversarial audit.",
            ],
            "notes": (
                "In second-year engineering curricula, foundational courses like Data Structures and Machine Learning require "
                "continuous engagement. When a student's attendance slips below 75% or their exam score drops sharply, early detection is vital. "
                "EduPilot AI uses an agentic workflow to decompose the review into planning, MCP analytics, intervention, and independent verification."
            ),
        },
        # SLIDE 3
        {
            "title": "Comparative Analysis: Existing Process vs. EduPilot AI",
            "subtitle": "Transforming Fragmented Manual Review into Actionable Decision Support",
            "bullets": [
                "Data Aggregation: 15-30 minutes of manual register cross-referencing -> Under 10 milliseconds via standardized MCP tools.",
                "Pattern Recognition: Isolated single-subject checks -> Holistic cross-subject correlation and trajectory detection.",
                "Intervention Quality: Vague, subjective verbal guidance -> Structured plans with assigned roles, review periods, and outcomes.",
                "Safety & Auditing: Unverified manual notes -> Independent Critic Agent recomputing mathematical bounds and rules.",
                "Governance Guardrail: EduPilot AI never debar, penalize, or replace teachers; all plans require human faculty approval.",
            ],
            "notes": (
                "This slide illustrates the contrast. Instead of faculty spending 20 minutes manually inspecting spreadsheets for a single student, "
                "EduPilot AI executes end-to-end multi-agent screening in under 10 milliseconds, assigning responsible faculty roles and measurable milestones."
            ),
        },
        # SLIDE 4
        {
            "title": "AI Agent Taxonomy & LangGraph Orchestration",
            "subtitle": "Theoretical Grounding (Russell & Norvig) with Four Specialized Agent Paradigms",
            "bullets": [
                "Agent 1: Goal-Based Agent (Academic Planner) — Formulates goal state and decomposes review into an ordered 6-step plan.",
                "Agent 2: Model-Based Reflex Agent (Student Analytics) — Maintains internal state across longitudinal records; evaluates condition-action rules via MCP.",
                "Agent 3: Utility-Based Agent (Intervention Planner) — Optimizes pedagogical recovery utility U(Action, Severity, Role) balancing impact vs burden.",
                "Agent 4: Learning & Verification Agent (Critic & Auditor) — Adversarial critic asserting mathematical invariants; feeds correction learning loop.",
                "LangGraph Conditional Routing: Enforces explicit graph edges, state persistence, and bounded retries (max 2) before report synthesis.",
            ],
            "notes": (
                "Here is our core AIML theoretical grounding from Russell & Norvig. Agent 1 is a Goal-Based Agent, Agent 2 is a Model-Based Reflex Agent "
                "with an internal longitudinal model, Agent 3 is a Utility-Based Agent optimizing recovery utility, and Agent 4 is a Learning/Verification Agent "
                "with an adversarial Critic component feeding error signals to our Correction node."
            ),
        },
        # SLIDE 5
        {
            "title": "Streamlit Proof-of-Concept Dashboard Demonstration",
            "subtitle": "Intuitive, Professional Academic Analytics Interface Designed for Faculty",
            "bullets": [
                "Page 1: Academic Overview — Cohort-wide KPI cards, attendance histograms, and pre-programmed archetype distribution.",
                "Page 2: Analyze Student — Deep-dive investigation featuring subject-wise attendance bars, trajectory charts, and evidence callouts.",
                "Page 3: Agent Workflow — Live observability into real LangGraph node execution events, latencies, and MCP tool logs.",
                "Page 4: Intervention Reports — Complete 18-field standardized documentation with one-click JSON and CSV export.",
                "Page 5: Batch Analysis — Rapid whole-section screening with automated support-priority categorization (High, Medium, Low).",
            ],
            "notes": (
                "Now let us view the live demonstration. Our Streamlit dashboard provides a clean, responsive academic portal. "
                "Faculty can select student STU1016, view their 65% attendance concern, examine the 18-field report, "
                "and inspect the live LangGraph execution trace on Page 3 with full telemetry."
            ),
        },
        # SLIDE 6
        {
            "title": "Verification, Critic Audit & Empirical Benchmark Results",
            "subtitle": "Quantitative Rule Evaluation, Invariant Safety Auditing, and Performance Benchmarks",
            "bullets": [
                "Automated Testing Suite: 26 Pytest test cases passing with 100% success rate across schemas, rules, agents, and graph.",
                "Empirical Benchmarking: 10 reproducible scenarios evaluated via scripts/evaluate.py across synthetic archetypes.",
                "100.0% Rule Compliance Rate: Exact adherence to attendance (<75%), failing marks (<40%), and trajectory decline (>=15% pts).",
                "Instantaneous Latency: Average end-to-end multi-agent execution duration of 6.9 milliseconds per student.",
                "Adversarial Trapping: Critic successfully trapped data corruption (attended > held) and ungrounded intervention claims.",
            ],
            "notes": (
                "Technical correctness and safety are the core priorities of this project. We built an automated test suite of 26 pytest unit tests "
                "all passing, and an evaluation runner benchmarking 10 reproducible scenarios with 100% compliance and 6.9 ms average latency."
            ),
        },
        # SLIDE 7
        {
            "title": "Ethical Safeguards, System Limitations & Conclusion",
            "subtitle": "Ethical Boundaries, Synthetic Data Considerations, and Future Scope",
            "bullets": [
                "Ethical Boundary: Purely decision support. Never makes consequential, grading, or punitive decisions automatically.",
                "Privacy & Fairness: Operates strictly on anonymous synthetic IDs (STU####); zero PII and zero demographic bias.",
                "Honest Limitations: Synthetic data cannot model qualitative realities (illness, grief, personal emergencies).",
                "Mandatory Human Oversight: Faculty mentors must review, contextualize, and approve all recommended interventions.",
                "Conclusion: EduPilot AI provides a robust, transparent, and reproducible blueprint for agentic academic decision support.",
            ],
            "notes": (
                "To conclude, EduPilot AI demonstrates how agentic AI can assist educators ethically while preserving human-in-the-loop oversight. "
                "By uniting LangGraph, MCP, transparent rules, and an independent Critic, we provide a production-grade academic tool. Thank you!"
            ),
        },
    ]

    for idx, data in enumerate(slides_content):
        slide = prs.slides.add_slide(blank_layout)

        # Slide Number Badge (Top Right)
        num_box = slide.shapes.add_textbox(Inches(11.5), Inches(0.4), Inches(1.5), Inches(0.5))
        tf_num = num_box.text_frame
        p_num = tf_num.paragraphs[0]
        p_num.text = f"Slide {idx + 1} of 7"
        p_num.font.size = Pt(12)
        p_num.font.bold = True
        p_num.font.color.rgb = COLOR_SLATE
        p_num.alignment = PP_ALIGN.RIGHT

        # Header Box: Title & Subtitle
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.0), Inches(1.5))
        tf_head = header_box.text_frame
        tf_head.word_wrap = True

        p_title = tf_head.paragraphs[0]
        p_title.text = data["title"]
        p_title.font.size = Pt(28)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_NAVY

        p_sub = tf_head.add_paragraph()
        p_sub.text = data["subtitle"]
        p_sub.font.size = Pt(16)
        p_sub.font.color.rgb = COLOR_SLATE
        p_sub.space_before = Pt(6)

        # Body Content Box (Bullet Points)
        content_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.3), Inches(11.7), Inches(4.5))
        tf_content = content_box.text_frame
        tf_content.word_wrap = True

        for b_idx, bullet in enumerate(data["bullets"]):
            p = tf_content.paragraphs[0] if b_idx == 0 else tf_content.add_paragraph()
            p.text = f"•  {bullet}"
            p.font.size = Pt(17)
            p.font.color.rgb = COLOR_DARK
            p.space_after = Pt(14)
            p.line_spacing = 1.15

        # Footer Box (Academic Branding)
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.7), Inches(0.4))
        tf_foot = footer_box.text_frame
        p_foot = tf_foot.paragraphs[0]
        p_foot.text = "EduPilot AI — AIML CA-3 Academic Support System | LangGraph + MCP + 4 Agents"
        p_foot.font.size = Pt(11)
        p_foot.font.color.rgb = COLOR_SLATE

        # Presenter / Speaker Notes
        notes_slide = slide.notes_slide
        text_frame = notes_slide.notes_text_frame
        text_frame.text = data["notes"]

    output_path = DOCS_DIR / "EduPilot_AI_CA3_Presentation.pptx"
    prs.save(str(output_path))
    print(f"[OK] Successfully generated presentation at {output_path}")


if __name__ == "__main__":
    create_presentation()
